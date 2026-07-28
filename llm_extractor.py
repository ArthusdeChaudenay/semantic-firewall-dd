import os
import re
import json
import time
from pathlib import Path

import fitz  # PyMuPDF
import pytesseract
from PIL import Image
from docx import Document as DocxDocument
from openai import OpenAI
from dotenv import load_dotenv

from benchmark_base import AUDITWEN_TEMPLATES, CHAMPS_SCHEMA

load_dotenv()

# --- Configuration Tesseract ---
TESSERACT_CMD = os.getenv("TESSERACT_CMD", r"C:\Program Files\Tesseract-OCR\tesseract.exe")
TESSDATA_PREFIX = os.getenv("TESSDATA_PREFIX", str(Path(__file__).parent / "tessdata"))
pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD
os.environ.setdefault("TESSDATA_PREFIX", TESSDATA_PREFIX)

# --- Configuration client OpenAI-compatible (Ollama, LM Studio, etc.) ---
# Valeurs par défaut pour Ollama local. Surcharger via .env si nécessaire.
_client = OpenAI(
    base_url=os.getenv("OPENAI_BASE_URL", "http://localhost:11434/v1"),
    api_key=os.getenv("OPENAI_API_KEY", "ollama"),
)
MODEL_NAME = os.getenv("MODEL_NAME", "qwen2.5:7b")

SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".png", ".jpg", ".jpeg", ".tiff", ".bmp", ".docx"}


# ==========================================
# EXTRACTION DE TEXTE MULTI-FORMAT
# ==========================================

def _extract_text_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8")


_RE_EURO_ARTIFACT = re.compile(r'([\d][\d\s,.]*)\s*\?(?=[\s\n,;)]|$)')

def _fix_pdf_encoding(text: str) -> str:
    """Remplace le symbole € corrompu en ? (artefact de polices PDF non-standard)."""
    return _RE_EURO_ARTIFACT.sub(r'\1 €', text)


def _extract_text_pdf(path: Path) -> str:
    """Extrait le texte d'un PDF ; bascule sur OCR pour les pages scannées."""
    doc = fitz.open(str(path))
    pages_text = []
    for page in doc:
        text = page.get_text().strip()
        if text:
            pages_text.append(_fix_pdf_encoding(text))
        else:
            pix = page.get_pixmap(dpi=300)
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            ocr_text = pytesseract.image_to_string(img, lang="fra").strip()
            pages_text.append(ocr_text)
    doc.close()
    return "\n".join(pages_text)


def _extract_text_image(path: Path) -> str:
    img = Image.open(str(path))
    return pytesseract.image_to_string(img, lang="fra").strip()


def _extract_text_docx(path: Path) -> str:
    doc = DocxDocument(str(path))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def extract_text_from_file(filepath: str) -> str:
    """Dispatch vers la bonne fonction d'extraction selon l'extension."""
    path = Path(filepath)
    ext = path.suffix.lower()

    if ext == ".txt":
        return _extract_text_txt(path)
    elif ext == ".pdf":
        return _extract_text_pdf(path)
    elif ext in (".png", ".jpg", ".jpeg", ".tiff", ".bmp"):
        return _extract_text_image(path)
    elif ext == ".docx":
        return _extract_text_docx(path)
    else:
        raise ValueError(f"Extension non supportée : '{ext}'. Extensions acceptées : {SUPPORTED_EXTENSIONS}")


# ==========================================
# APPEL QWEN VIA API OPENAI-COMPATIBLE
# ==========================================

def _build_extraction_prompt(document_text: str) -> str:
    schema_str = json.dumps(CHAMPS_SCHEMA, indent=4, ensure_ascii=False)
    return AUDITWEN_TEMPLATES["extraction_structuree"].format(
        schema=schema_str,
        document_text=document_text,
    )


def _clean_llm_response(raw: str) -> str:
    """Supprime les éventuelles balises Markdown que le modèle aurait ajoutées."""
    raw = raw.strip()
    if raw.startswith("```"):
        # Retire la première ligne (```json ou ```) et la dernière (```)
        lines = raw.split("\n")
        raw = "\n".join(lines[1:]).rsplit("```", 1)[0].strip()
    return raw


_MAX_RETRIES  = 4
_RETRY_BASE_S = 2.0   # attente initiale en secondes (doublée à chaque tentative)


def _llm_call_with_retry(messages: list, max_retries: int = _MAX_RETRIES) -> str:
    """
    Appel LLM avec retry exponentiel sur 429 (Too Many Requests).
    Délais : 2s → 4s → 8s → 16s avant abandon.
    """
    import openai

    delay = _RETRY_BASE_S
    for attempt in range(max_retries + 1):
        try:
            response = _client.chat.completions.create(
                model=MODEL_NAME,
                messages=messages,
                temperature=0,
                max_tokens=4096,
            )
            return response.choices[0].message.content
        except openai.RateLimitError as exc:
            if attempt == max_retries:
                raise
            print(f"    [LLM] 429 rate-limit — attente {delay:.0f}s (tentative {attempt+1}/{max_retries})…")
            time.sleep(delay)
            delay *= 2
        except openai.APIStatusError as exc:
            if exc.status_code == 429:
                if attempt == max_retries:
                    raise
                print(f"    [LLM] 429 rate-limit — attente {delay:.0f}s (tentative {attempt+1}/{max_retries})…")
                time.sleep(delay)
                delay *= 2
            else:
                raise


def call_qwen_extraction(document_text: str) -> dict:
    """
    Envoie le texte du document à Qwen avec le template AuditWen et retourne
    le dictionnaire de champs extraits.
    Lève json.JSONDecodeError si le modèle ne produit pas un JSON valide.
    """
    prompt = _build_extraction_prompt(document_text)
    raw_output = _llm_call_with_retry([{"role": "user", "content": prompt}])
    cleaned = _clean_llm_response(raw_output)
    return json.loads(cleaned)


# ==========================================
# FONCTION PRINCIPALE D'EXTRACTION
# ==========================================

def extract_document(filepath: str) -> dict:
    """
    Pipeline complet pour un document :
      1. Extraction du texte brut (txt / pdf / image / docx)
      2. Appel Qwen avec template AuditWen
      3. Mise en forme dans le schéma de référence FinVerBench

    Retourne toujours un dict avec les clés :
        nom_fichier, statut ("ok" | "error"), champs, message
    """
    nom_fichier = Path(filepath).name

    try:
        text = extract_text_from_file(filepath)

        if not text.strip():
            return {
                "nom_fichier": nom_fichier,
                "statut": "error",
                "champs": {},
                "message": "Aucun texte extrait du document.",
            }

        extracted = call_qwen_extraction(text)

        champs = {}
        for key, val in extracted.items():
            if key.endswith("_brut"):
                continue  # rattaché au champ principal ci-dessous
            champs[key] = {"valeur": str(val) if val is not None else None, "source": "llm"}
            brut_val = extracted.get(f"{key}_brut")
            if brut_val is not None:
                champs[key]["valeur_brute"] = str(brut_val)

        # Fallback déterministe : date_echeance absente → chercher les alias dans le texte brut
        if not champs.get("date_echeance", {}).get("valeur"):
            _m = re.search(
                r'(?:échéance(?:\s+de\s+paiement)?|date\s+limite(?:\s+de\s+paiement)?)'
                r'\s*[:\-]?\s*(\d{2}/\d{2}/\d{4})',
                text, re.IGNORECASE
            )
            if _m:
                champs.setdefault("date_echeance", {})
                champs["date_echeance"]["valeur"] = _m.group(1)
                champs["date_echeance"]["source"] = "regex_fallback"

        return {"nom_fichier": nom_fichier, "statut": "ok", "champs": champs, "message": None}

    except json.JSONDecodeError as e:
        return {
            "nom_fichier": nom_fichier,
            "statut": "error",
            "champs": {},
            "message": f"Réponse LLM non parseable en JSON : {e}",
        }
    except Exception as e:
        return {
            "nom_fichier": nom_fichier,
            "statut": "error",
            "champs": {},
            "message": str(e),
        }


# ==========================================
# PASS 2 AUDITWEN : VÉRIFICATION LLM
# ==========================================

def call_qwen_verification(flat_data: dict) -> dict:
    """
    Pass 2 AuditWen : demande à Qwen de contrôler la cohérence de sa propre extraction.
    flat_data = dict {champ: valeur} sans la couche {"valeur": ..., "source": ...}.
    Lève json.JSONDecodeError si la réponse n'est pas un JSON valide.
    """
    prompt = AUDITWEN_TEMPLATES["verification_coherence"].format(
        extracted_json=json.dumps(flat_data, indent=2, ensure_ascii=False)
    )
    raw = _llm_call_with_retry([{"role": "user", "content": prompt}])
    return json.loads(_clean_llm_response(raw))


# ==========================================
# ZONE DE TEST RAPIDE
# ==========================================
if __name__ == "__main__":
    import sys
    target = sys.argv[1] if len(sys.argv) > 1 else "samples/facture_exemple.txt"
    result = extract_document(target)
    print(json.dumps(result, indent=4, ensure_ascii=False))
